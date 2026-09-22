"""One authorized data model and A4 layout for browser, PDF and editable PowerPoint."""
import io
import math
import os
from collections import defaultdict
from datetime import date, timedelta
from html import escape
from pathlib import Path

from catalog import DEPARTMENTS
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

WIDTH, HEIGHT, MARGIN = 595.276, 841.89, 34.016
BLUE, GREEN, GRAY, INK = '#254d3c', '#7d9c67', '#edf2e7', '#243c35'
FONT_FILE = os.getenv('UAG_FONT') or next((str(p) for p in [Path('C:/Windows/Fonts/arial.ttf'), Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf')] if p.exists()), None)
LOGO = Path(__file__).resolve().parent / 'static' / 'logo.png'
if not FONT_FILE:
    raise RuntimeError('Кирилл фонт олдсонгүй. UAG_FONT орчны хувьсагчид TTF зам өгнө үү.')
pdfmetrics.registerFont(TTFont('UAG', FONT_FILE))


class ReportOverflow(ValueError):
    pass


def progress(value):
    try:
        number = float(str(value).strip().rstrip('%'))
        return number if 0 <= number <= 100 else None
    except (TypeError, ValueError):
        return None


def metrics(rows):
    values = [progress(r['progress']) for r in rows]
    known = [p for p in values if p is not None]
    total = len(rows)
    complete = sum(p == 100 for p in known)
    ongoing = sum(0 < p < 100 for p in known)
    return dict(total=total, complete=complete, ongoing=ongoing, not_started=sum(p == 0 for p in known),
                unknown=total-len(known), complete_pct=round(complete / total * 100, 1) if total else 0,
                ongoing_pct=round(ongoing / total * 100, 1) if total else 0,
                average=round(sum(known) / len(known), 1) if known else None,
                issues=sum(bool((r.get('issue') or '').strip()) for r in rows))


def wrap(text, width, size):
    """Measured Unicode wrapping; explicit newlines and long tokens are preserved."""
    lines = []
    for paragraph in str(text or '').split('\n'):
        current = ''
        for word in paragraph.split(' '):
            candidate = (current + ' ' + word).strip()
            if pdfmetrics.stringWidth(candidate, 'UAG', size) <= width:
                current = candidate
                continue
            if current: lines.append(current)
            current = ''
            for char in word:
                if current and pdfmetrics.stringWidth(current + char, 'UAG', size) > width:
                    lines.append(current); current = ''
                current += char
        lines.append(current)
    return lines or ['']


def label(page, text, x, y, w, size=11, color=INK):
    lines = wrap(text, w, size)
    block = dict(kind='text', text=str(text), lines=lines, x=x, y=y, w=w, h=len(lines)*size*1.3+4, size=size, color=color)
    page['blocks'].append(block)
    return y + block['h']


def table(page, headers, rows, widths, y, size=9.5):
    cells, heights = [], []
    for row in [headers] + rows:
        wrapped = [wrap(value, width-12, size) for value, width in zip(row, widths)]
        cells.append(wrapped)
        heights.append(max(len(lines) for lines in wrapped)*size*1.3+12)
    block = dict(kind='table', headers=headers, rows=rows, cells=cells, heights=heights, widths=widths,
                 x=MARGIN, y=y, w=sum(widths), h=sum(heights), size=size)
    page['blocks'].append(block)
    return y + block['h']


def chart(page, kind, title, labels, values, x, y, w, h):
    block = dict(kind='chart', chart=kind, title=title, labels=labels, values=values, x=x, y=y, w=w, h=h)
    block['svg'] = chart_svg(block)
    page['blocks'].append(block)


def chart_svg(b):
    w, h = b['w'], b['h']
    parts = [f'<svg viewBox="0 0 {w} {h}" xmlns="http://www.w3.org/2000/svg" role="img" aria-label="{escape(b["title"], quote=True)}">']
    def text(x, y, value, size=9):
        parts.append(f'<text x="{x}" y="{y}" fill="{INK}" font-size="{size}" font-family="Arial">{escape(str(value))}</text>')
    text(0, 14, b['title'], 11)
    values, labels = b['values'], b['labels']
    if b['chart'] == 'pie':
        total = sum(values)
        palette = [GREEN, BLUE, '#91a8bc', '#b88449']
        cx, cy, radius = w/2, 90, 57
        angle = -math.pi/2
        if not total:
            parts.append(f'<circle cx="{cx}" cy="{cy}" r="{radius}" fill="{GRAY}"/>')
            text(20, 166, 'Өгөгдөл алга')
        for i, value in enumerate(values):
            if total and value:
                end = angle+value/total*math.tau
                if value == total:
                    parts.append(f'<circle cx="{cx}" cy="{cy}" r="{radius}" fill="{palette[i]}"/>')
                else:
                    a,bx = cx+radius*math.cos(angle), cy+radius*math.sin(angle)
                    c,d = cx+radius*math.cos(end), cy+radius*math.sin(end)
                    parts.append(f'<path d="M {cx} {cy} L {a} {bx} A {radius} {radius} 0 {int(value/total>.5)} 1 {c} {d} Z" fill="{palette[i]}"/>')
                angle=end
            text(5, 174+i*17, f'{labels[i]}: {value}')
    elif b['chart'] == 'bar':
        for i, (name, value) in enumerate(zip(labels, values)):
            y = 42+i*40
            text(0,y-8,name)
            parts.append(f'<rect x="0" y="{y}" width="{w-38}" height="9" fill="{GRAY}"/>')
            if value is not None:
                parts.append(f'<rect x="0" y="{y}" width="{(w-38)*value/100}" height="9" fill="{BLUE}"/>')
            text(w-35,y+9,'—' if value is None else f'{value:g}%')
    else:
        left, top, cw, ch = 28, 33, w-48, h-69
        for tick in (0,50,100):
            y=top+ch*(1-tick/100)
            text(0,y+3,str(tick),8)
            parts.append(f'<path d="M {left} {y} H {left+cw}" stroke="{GRAY}"/>')
        previous=None
        for i,(name,value) in enumerate(zip(labels,values)):
            x=left+i*cw/max(1,len(values)-1)
            text(x-12,top+ch+18,name,8)
            if value is None:
                previous=None
                continue
            y=top+ch*(1-value/100)
            if previous: parts.append(f'<path d="M {previous[0]} {previous[1]} L {x} {y}" stroke="{BLUE}" stroke-width="2"/>')
            parts.append(f'<circle cx="{x}" cy="{y}" r="3" fill="{GREEN}"/>')
            text(x-8,y-7,f'{value:g}',8)
            previous=(x,y)
    return ''.join(parts)+'</svg>'


def build_report(rows, history, dept, start, end, leader):
    selected = [dept] if dept else DEPARTMENTS
    summary = metrics(rows)
    model = dict(pages=[], start=start, end=end, metrics=summary, overflow=[], unknown_dates=sum(not r.get('report_day') for r in history))
    usable = WIDTH - 2*MARGIN

    def new_page(title):
        page = dict(title=title, blocks=[])
        model['pages'].append(page)
        if LOGO.exists():
            page['blocks'].append(dict(kind='image', x=WIDTH-MARGIN-80, y=16, w=80, h=39.76, src='logo.png'))
        label(page, 'ҮЙЛ АЖИЛЛАГААНЫ ГАЗАР', MARGIN, 31, usable, 11, BLUE)
        label(page, title, MARGIN, 58, usable, 21)
        label(page, f'Тайлант хугацаа: {start or "эхнээс"} – {end or "өнөөдөр"}', MARGIN, 92, usable, 10)
        label(page, f'ЦЭЦЭНС МАЙНИНГ ЭНД ЭНЕРЖИ ХХК     •     {len(model["pages"])}', MARGIN, HEIGHT-29, usable, 8, BLUE)
        return page

    def kpis(page, values, y=126):
        for i,(name,value) in enumerate(values):
            x=MARGIN+i*usable/4
            label(page, value, x, y, usable/4-8, 23, BLUE)
            label(page, name, x, y+34, usable/4-8, 9)

    if leader:
        p=new_page('Удирдлагын нэгдсэн тайлан')
        kpis(p,[('Нийт ажил',str(summary['total'])),('Дууссан',f'{summary["complete_pct"]}%'),('Явцтай',f'{summary["ongoing_pct"]}%'),('Шийдвэрлэх асуудал',str(summary['issues']))])
        chart(p,'pie','Ажлын төлөв', ['Дууссан','Явцтай','Эхлээгүй','Явц тодорхойгүй'],
              [summary[k] for k in ('complete','ongoing','not_started','unknown')],MARGIN,200,220,252)
        chart(p,'bar','Хэлтсийн дундаж гүйцэтгэл',selected,[metrics([r for r in rows if r['department']==d])['average'] for d in selected],MARGIN+250,200,usable-250,252)
        table(p,['Хэлтэс','Ажил','Дууссан','Дундаж явц'],[[d,str((m:=metrics([r for r in rows if r['department']==d]))['total']),str(m['complete']), '—' if m['average'] is None else f'{m["average"]}%'] for d in selected],[usable*.49,usable*.13,usable*.16,usable*.22],485)
        label(p,'Дууссан: явц 100%. Явцтай: 0%-иас их, 100%-иас бага. Хувийн суурь нь бүх ажил. Дундаж явцад зөвхөн тоон явцтай ажлууд орно.', MARGIN,710,usable,9)
        p=new_page('Гүйцэтгэлийн чиг хандлага, асуудал')
        last=date.fromisoformat(end) if end else date.today()
        last-=timedelta(days=last.weekday())
        weeks=[last-timedelta(weeks=i) for i in reversed(range(8))]
        vals=[]
        for week in weeks:
            group=[r for r in history if r.get('report_day') and week.isoformat()<=r['report_day']<=(week+timedelta(days=6)).isoformat()]
            vals.append(metrics(group)['average'])
        chart(p,'line','Сүүлийн 8 долоо хоногийн дундаж явц (%)',[w.strftime('%m.%d') for w in weeks],vals,MARGIN,126,usable,190)
        label(p,'Өгөгдөлгүй долоо хоногт цэг тавихгүй. Тренд нь сонгосон хэлтэс, дуусах огноонд хамаарна.',MARGIN,325,usable,9)
        label(p,'Шийдвэрлүүлэх асуудал',MARGIN,360,usable,14)
        issues=[[r['department'],r.get('issue',''),r.get('assignee') or '—'] for r in rows if (r.get('issue') or '').strip()]
        table(p,['Хэлтэс','Асуудал','Хариуцагч'],issues or [['—','Бүртгэсэн асуудал байхгүй','—']],[usable*.23,usable*.57,usable*.20],392)
    for d in selected:
        p=new_page(d)
        group=[r for r in rows if r['department']==d]
        m=metrics(group)
        kpis(p,[('Нийт ажил',str(m['total'])),('Дууссан',f'{m["complete_pct"]}%'),('Дундаж явц','—' if m['average'] is None else f'{m["average"]}%'),('Шийдвэрлэх асуудал',str(m['issues']))])
        names=list(dict.fromkeys(r['func_area'] for r in group))
        if names:
            # Compact native chart, keeping space for the full task content.
            chart(p,'bar','Чиг үүргийн гүйцэтгэл',names,[metrics([r for r in group if r['func_area']==n])['average'] for n in names],MARGIN,198,usable,40*len(names)+30)
        y=242+40*len(names) if names else 205
        body=[]
        for r in group:
            detail=r['task_desc']
            if r.get('task_result'): detail+='\nҮр дүн: '+r['task_result']
            if r.get('issue'): detail+='\nШийдвэрлүүлэх: '+r['issue']
            body.append([r['func_area'],detail,r['progress'],r.get('assignee') or '—'])
        table(p,['Чиг үүрэг','Ажил, үр дүн, асуудал','Явц','Хариуцагч'],body or [['—','Тайлант хугацаанд бүртгэл алга','—','—']],[usable*.20,usable*.55,usable*.09,usable*.16],y)
    for p in model['pages']:
        for b in p['blocks']:
            if b['y'] < HEIGHT-40 and b['y']+b['h'] > HEIGHT-48:
                model['overflow'].append(p['title']); break
    return model


def ensure_fit(model):
    if model['overflow']:
        raise ReportOverflow('Нэг A4 хуудсанд багтахгүй тайлан: '+', '.join(model['overflow'])+'. Тайлант хугацааг багасгах эсвэл ажлын тайлбарыг товчилно уу. Бүх дэлгэрэнгүй мэдээллийг CSV-ээр татаж болно.')


def export_pdf(model):
    from reportlab.pdfgen import canvas
    from reportlab.lib.colors import HexColor
    ensure_fit(model)
    output=io.BytesIO(); c=canvas.Canvas(output,pagesize=(WIDTH,HEIGHT))
    c.setTitle('ҮАГ 7 хоногийн тайлан')
    def lines(texts,x,y,size,color=INK):
        c.setFillColor(HexColor(color)); c.setFont('UAG',size)
        for i,text in enumerate(texts): c.drawString(x,HEIGHT-y-size-i*size*1.3,text)
    for page in model['pages']:
        for b in page['blocks']:
            if b['kind']=='image': c.drawImage(str(LOGO),b['x'],HEIGHT-b['y']-b['h'],b['w'],b['h'],mask='auto')
            elif b['kind']=='text': lines(b['lines'],b['x'],b['y'],b['size'],b['color'])
            elif b['kind']=='table':
                y=b['y']
                for row_idx,(cells,height) in enumerate(zip(b['cells'],b['heights'])):
                    x=b['x']
                    c.setFillColor(HexColor(GRAY if row_idx==0 else '#ffffff'))
                    c.rect(x,HEIGHT-y-height,b['w'],height,fill=1,stroke=0)
                    for cell,width in zip(cells,b['widths']):
                        lines(cell,x+6,y+6,b['size']); x+=width
                    c.setStrokeColor(HexColor(GRAY)); c.line(b['x'],HEIGHT-y-height,b['x']+b['w'],HEIGHT-y-height)
                    y+=height
            else:
                x,y,w,h=b['x'],b['y'],b['w'],b['h']
                lines([b['title']],x,y+2,11)
                if b['chart']=='pie':
                    total=sum(b['values']); angle=90
                    if not total:
                        c.setFillColor(HexColor(GRAY)); c.circle(x+w/2,HEIGHT-y-90,57,fill=1,stroke=0)
                    for i,(name,value) in enumerate(zip(b['labels'],b['values'])):
                        if total and value:
                            c.setFillColor(HexColor([GREEN,BLUE,'#91a8bc','#b88449'][i]))
                            extent=-360*value/total
                            c.wedge(x+w/2-57,HEIGHT-y-147,x+w/2+57,HEIGHT-y-33,angle,extent,fill=1,stroke=0)
                            angle+=extent
                        lines([f'{name}: {value}'],x+5,y+164+i*17,9)
                elif b['chart']=='bar':
                    for i,(name,value) in enumerate(zip(b['labels'],b['values'])):
                        yy=y+42+i*40
                        lines([name],x,yy-18,9)
                        c.setFillColor(HexColor(GRAY)); c.rect(x,HEIGHT-yy-9,w-38,9,fill=1,stroke=0)
                        if value is not None:
                            c.setFillColor(HexColor(BLUE)); c.rect(x,HEIGHT-yy-9,(w-38)*value/100,9,fill=1,stroke=0)
                        lines(['—' if value is None else f'{value:g}%'],x+w-35,yy,9)
                else:
                    left,top,cw,ch=x+28,y+33,w-48,h-69
                    for tick in (0,50,100):
                        yy=top+ch*(1-tick/100)
                        lines([str(tick)],x,yy-5,8)
                        c.setStrokeColor(HexColor(GRAY)); c.line(left,HEIGHT-yy,left+cw,HEIGHT-yy)
                    prev=None
                    for i,(name,value) in enumerate(zip(b['labels'],b['values'])):
                        xx=left+i*cw/max(1,len(b['values'])-1)
                        lines([name],xx-12,top+ch+8,8)
                        if value is None: prev=None; continue
                        yy=top+ch*(1-value/100)
                        if prev:
                            c.setStrokeColor(HexColor(BLUE)); c.setLineWidth(1.5); c.line(prev[0],HEIGHT-prev[1],xx,HEIGHT-yy)
                        c.setFillColor(HexColor(GREEN)); c.circle(xx,HEIGHT-yy,3,stroke=0,fill=1)
                        lines([f'{value:g}'],xx-8,yy-18,8); prev=(xx,yy)
        c.showPage()
    c.save(); return output.getvalue()


def export_pptx(model):
    # The user explicitly requested python-pptx as the application export engine.
    from pptx import Presentation
    from pptx.util import Pt
    from pptx.dml.color import RGBColor
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
    ensure_fit(model)
    deck=Presentation(); deck.slide_width=Pt(WIDTH); deck.slide_height=Pt(HEIGHT)
    def text_box(slide,lines,x,y,w,h,size,color=INK):
        box=slide.shapes.add_textbox(Pt(x),Pt(y),Pt(w),Pt(h+6))
        frame=box.text_frame; frame.word_wrap=False
        frame.margin_left=frame.margin_right=frame.margin_top=frame.margin_bottom=0
        for i,line in enumerate(lines):
            p=frame.paragraphs[0] if i==0 else frame.add_paragraph()
            p.text=line; p.font.name='Arial'; p.font.size=Pt(size); p.font.color.rgb=RGBColor.from_string(color[1:]); p.space_after=Pt(0); p.space_before=Pt(0); p.line_spacing=1.3
    for page in model['pages']:
        slide=deck.slides.add_slide(deck.slide_layouts[6])
        for b in page['blocks']:
            if b['kind']=='image': slide.shapes.add_picture(str(LOGO),Pt(b['x']),Pt(b['y']),width=Pt(b['w']),height=Pt(b['h']))
            elif b['kind']=='text': text_box(slide,b['lines'],b['x'],b['y'],b['w'],b['h'],b['size'],b['color'])
            elif b['kind']=='table':
                tab=slide.shapes.add_table(len(b['cells']),len(b['widths']),Pt(b['x']),Pt(b['y']),Pt(b['w']),Pt(b['h'])).table
                for col,width in zip(tab.columns,b['widths']): col.width=Pt(width)
                for i,(cells,height) in enumerate(zip(b['cells'],b['heights'])):
                    tab.rows[i].height=Pt(height)
                    for j,lines in enumerate(cells):
                        cell=tab.cell(i,j); cell.text='\n'.join(lines)
                        cell.margin_left=cell.margin_right=Pt(6); cell.margin_top=cell.margin_bottom=Pt(5)
                        cell.fill.solid(); cell.fill.fore_color.rgb=RGBColor.from_string((GRAY if i==0 else '#ffffff')[1:])
                        for p in cell.text_frame.paragraphs:
                            p.font.name='Arial'; p.font.size=Pt(b['size']); p.font.color.rgb=RGBColor.from_string(INK[1:]); p.line_spacing=1.3; p.space_after=Pt(0)
            else:
                text_box(slide,[b['title']],b['x'],b['y'],b['w'],22,11)
                if not any(v is not None and v != 0 for v in b['values']):
                    text_box(slide,['Өгөгдөл алга' if all(v is None for v in b['values']) or b['chart']=='pie' else 'Гүйцэтгэл: 0%'],b['x'],b['y']+40,b['w'],30,11)
                    continue
                data=CategoryChartData(); data.categories=b['labels']; data.add_series('Гүйцэтгэл',b['values'])
                kind={'pie':XL_CHART_TYPE.PIE,'bar':XL_CHART_TYPE.BAR_CLUSTERED,'line':XL_CHART_TYPE.LINE_MARKERS}[b['chart']]
                chart=slide.shapes.add_chart(kind,Pt(b['x']),Pt(b['y']+24),Pt(b['w']),Pt(b['h']-24),data).chart
                chart.font.name='Arial'; chart.font.size=Pt(9); chart.has_legend=b['chart']=='pie'
                if chart.has_legend:
                    chart.legend.position=XL_LEGEND_POSITION.BOTTOM; chart.legend.include_in_layout=False
                if b['chart']!='pie':
                    chart.value_axis.minimum_scale=0; chart.value_axis.maximum_scale=100
                    chart.value_axis.tick_labels.font.size=Pt(8); chart.category_axis.tick_labels.font.size=Pt(8)
                chart.plots[0].has_data_labels=True
                chart.plots[0].data_labels.font.size=Pt(8)
        slide.notes_slide.notes_text_frame.text='Тайлант огноогоор шүүсэн системийн өгөгдөл. Дууссан = 100%. Явцтай = 0 < явц < 100. Дундаж явцад тоон утгатай ажлууд орно.'
    output=io.BytesIO(); deck.save(output); return output.getvalue()
