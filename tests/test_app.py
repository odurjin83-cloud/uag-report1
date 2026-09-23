import io
import sqlite3
import tempfile
import unittest
from contextlib import contextmanager
from pathlib import Path
from urllib.parse import urlencode

from pptx import Presentation
from werkzeug.security import generate_password_hash
from werkzeug.datastructures import MultiDict
from app import create_app, init_db
from catalog import FUNCTIONS, DEPARTMENTS
from reporting import build_report, export_pdf, export_pptx, metrics, ReportOverflow


@contextmanager
def connection(path):
    db=sqlite3.connect(path)
    try:
        with db: yield db
    finally: db.close()


class ApplicationTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory()
        self.path=Path(self.tmp.name)/'test.db'
        self.app=create_app({'TESTING':True,'DATABASE':str(self.path),'SECRET_KEY':'test-key'})
        self.client=self.app.test_client()
        with connection(self.path) as db:
            for email,role,dept in [('it@example.com','user','МТХэлтэс'),('boss@example.com','admin','Бүх хэлтэс'),('manager@example.com','manager','МТХэлтэс'),('lead@example.com','management','Бүх хэлтэс')]:
                db.execute('INSERT INTO users(email,password,role,department) VALUES (?,?,?,?)',(email,generate_password_hash('test-password'),role,dept))
            for i,dept in enumerate(DEPARTMENTS):
                db.execute('INSERT INTO reports(department,report_date,report_day,func_area,task_desc,task_result,progress,assignee,issue,created_at) VALUES (?,?,?,?,?,?,?,?,?,?)',
                           (dept,'2026.09.20 - 38-р 7 хоног','2026-09-20',FUNCTIONS[dept][0],f'private-{i}','Үр дүн',f'{i*25}%','Бат','Шийдвэр гаргах' if i==3 else '', '2026-09-22T13:00:00'))
        self.query='?start_date=2026-09-14&end_date=2026-09-20'

    def tearDown(self): self.tmp.cleanup()

    def login(self,email='it@example.com'):
        self.client.get('/login')
        with self.client.session_transaction() as session: token=session['csrf_token']
        response=self.client.post('/login',data={'csrf_token':token,'email':email,'password':'test-password'})
        self.assertEqual(response.status_code,302)
        self.client.get('/')

    def post(self,url,data=None):
        with self.client.session_transaction() as session: token=session['csrf_token']
        payload=MultiDict(data or {}); payload.add('csrf_token',token)
        return self.client.post(url,data=payload)

    def report_form(self):
        return dict(department='ЗОНХХэлтэс',report_date='2026-09-20',func_area='ERP / Help desk',task_desc='Систем шалгах',task_result='Шалгасан',progress='100',assignee='Бат',issue='')

    def test_login_required(self):
        for url in ['/','/reports','/reports/export/pdf','/reports/export/pptx','/export_csv']:
            self.assertEqual(self.client.get(url).status_code,302)

    def test_user_scope_every_read_and_export(self):
        self.login()
        foreign='&dept='+DEPARTMENTS[0]
        for url in ['/'+self.query+foreign,'/reports'+self.query+foreign,'/analytics'+self.query+foreign,'/export_csv'+self.query+foreign]:
            response=self.client.get(url)
            self.assertEqual(response.status_code,200)
            self.assertIn(b'private-3',response.data)
            self.assertNotIn(b'private-0',response.data)
        html=self.client.get('/').get_data(as_text=True)
        self.assertIn('ERP / Help desk',html)
        self.assertNotIn('PR, олон нийт',html)
        for kind in ['pdf','pptx']:
            response=self.client.get('/reports/export/'+kind+self.query+foreign)
            self.assertEqual(response.status_code,200)
            if kind=='pptx':
                deck=Presentation(io.BytesIO(response.data))
                self.assertEqual(len(deck.slides),1)
                texts=' '.join(sh.text for sl in deck.slides for sh in sl.shapes if sh.has_text_frame)
                self.assertNotIn(DEPARTMENTS[0],texts)
            else: self.assertTrue(response.data.startswith(b'%PDF'))

    def test_forged_category_atomic_batch_and_department(self):
        self.login()
        form=self.report_form()
        form['func_area']='PR, олон нийт'
        self.assertEqual(self.post('/save',form).status_code,400)
        form=self.report_form()
        self.assertEqual(self.post('/save',form).status_code,302)
        with connection(self.path) as db:
            self.assertEqual(db.execute('SELECT department FROM reports ORDER BY id DESC').fetchone()[0],'МТХэлтэс')
        data=MultiDict(self.report_form()); data.add('func_area','PR, олон нийт')
        for k in ('task_desc','task_result','progress','assignee','issue'): data.add(k,self.report_form()[k])
        self.assertEqual(self.post('/save',data).status_code,400)
        with connection(self.path) as db: self.assertEqual(db.execute('SELECT COUNT(*) FROM reports').fetchone()[0],6)

    def test_edit_ownership_and_locked_reports(self):
        self.login()
        self.assertEqual(self.client.get('/edit_report/1').status_code,404)
        self.assertEqual(self.post('/edit_report/1',self.report_form()).status_code,404)
        with connection(self.path) as db: db.execute("UPDATE reports SET status='approved' WHERE id=4")
        self.assertEqual(self.post('/edit_report/4',self.report_form()).status_code,409)

    def test_roles_admin_pages_and_filter(self):
        for email in ['boss@example.com','lead@example.com']:
            self.login(email)
            response=self.client.get('/reports/export/pptx'+self.query)
            self.assertEqual(response.status_code,200)
            self.assertEqual(len(Presentation(io.BytesIO(response.data)).slides),7)
            response=self.client.get('/reports/export/pptx'+self.query+'&'+urlencode({'dept':'МТХэлтэс'}))
            self.assertEqual(len(Presentation(io.BytesIO(response.data)).slides),3)
        self.login('manager@example.com')
        self.assertEqual(self.client.get('/admin').status_code,403)
        response=self.client.get('/reports/export/pptx'+self.query)
        self.assertEqual(len(Presentation(io.BytesIO(response.data)).slides),1)

    def test_database_role_revocation_applies_immediately(self):
        self.login('boss@example.com')
        with connection(self.path) as db: db.execute("UPDATE users SET role='user',department='МТХэлтэс' WHERE email='boss@example.com'")
        self.assertEqual(self.client.get('/admin').status_code,403)
        self.assertNotIn(b'private-0',self.client.get('/').data)

    def test_report_date_not_creation_and_bad_dates(self):
        self.login()
        self.assertIn(b'private-3',self.client.get('/reports'+self.query).data)
        for query in ['?start_date=bad','?start_date=2026-09-22&end_date=2026-09-01']:
            self.assertEqual(self.client.get('/reports'+query).status_code,400)

    def test_csrf_and_escaped_content(self):
        self.login()
        self.assertEqual(self.client.post('/save',data=self.report_form()).status_code,400)
        form=self.report_form(); form['task_desc']='<script>alert(1)</script>{{7*7}}'
        self.assertEqual(self.post('/save',form).status_code,302)
        html=self.client.get('/').get_data(as_text=True)
        self.assertIn('&lt;script&gt;',html)
        self.assertIn('{{7*7}}',html)
        self.assertNotIn('<script>alert(1)</script>',html)

    def test_missing_department_is_denied(self):
        self.login()
        with connection(self.path) as db: db.execute("UPDATE users SET department=NULL WHERE email='it@example.com'")
        self.assertEqual(self.client.get('/').status_code,403)


class MigrationTests(unittest.TestCase):
    def test_migration_preserves_legacy_rows_and_is_idempotent(self):
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'legacy.db'
            with connection(path) as db:
                db.execute('CREATE TABLE categories(id INTEGER PRIMARY KEY,name TEXT UNIQUE NOT NULL)')
                db.executemany('INSERT INTO categories VALUES (?,?)',[(1,'PR, олон нийт'),(2,'Төв кемп'),(3,'Unknown category')])
                db.execute('CREATE TABLE reports(id INTEGER PRIMARY KEY, department TEXT, report_date TEXT, func_area TEXT, task_desc TEXT, task_result TEXT, progress TEXT, assignee TEXT, status TEXT, created_at TEXT)')
                db.execute("INSERT INTO reports VALUES(1,'МТХэлтэс','2026.09.20 - 38-р 7 хоног','PR, олон нийт','Old task','','50%','','submitted','2026-09-22')")
            init_db(path)
            with connection(path) as db:
                first=db.execute('SELECT COUNT(*) FROM categories').fetchone()[0]
                self.assertEqual(db.execute('SELECT report_day,task_desc FROM reports').fetchone(),('2026-09-20','Old task'))
                self.assertEqual(db.execute('SELECT department FROM categories WHERE id=1').fetchone()[0],'ЗОНХХэлтэс')
                self.assertIsNone(db.execute('SELECT department FROM categories WHERE id=3').fetchone()[0])
                db.execute('INSERT INTO categories(name,department) VALUES (?,?)',('PR, олон нийт','МТХэлтэс'))
            init_db(path)
            with connection(path) as db: self.assertEqual(db.execute('SELECT COUNT(*) FROM categories').fetchone()[0],first+1)
            self.assertEqual(len(list(Path(tmp).glob('*.before-v2-*'))),1)

    def test_empty_metrics_and_overflow(self):
        model=build_report([],[],'МТХэлтэс','2026-09-14','2026-09-20',False)
        self.assertEqual(model['metrics']['total'],0)
        self.assertTrue(export_pdf(model).startswith(b'%PDF'))
        self.assertEqual(len(Presentation(io.BytesIO(export_pptx(model))).slides),1)
        row=dict(department='МТХэлтэс',func_area='ERP / Help desk',task_desc='Урт тайлбар '*2000,progress='50%',task_result='',issue='',assignee='Бат',report_day='2026-09-20')
        model=build_report([row],[row],'МТХэлтэс','2026-09-14','2026-09-20',False)
        self.assertTrue(model['overflow'])
        for exporter in (export_pdf,export_pptx):
            with self.assertRaises(ReportOverflow): exporter(model)
        m=metrics([dict(progress=p,issue='') for p in ['100%','50%','0%','Unknown']])
        self.assertEqual((m['complete_pct'],m['ongoing_pct'],m['unknown'],m['average']),(25,25,1,50))


if __name__=='__main__': unittest.main()
