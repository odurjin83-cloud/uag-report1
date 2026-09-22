"""UAG weekly reports. Run python app.py, or waitress-serve --call app:create_app."""
import csv
import io
import os
import re
import secrets
import sqlite3
from datetime import date, datetime, timedelta
from functools import wraps
from contextlib import closing
from pathlib import Path

from flask import Flask, abort, g, redirect, render_template, request, send_file, session, url_for
from werkzeug.security import check_password_hash, generate_password_hash
from catalog import DEPARTMENTS, FUNCTIONS, LEADERS, ROLES

ROOT = Path(__file__).resolve().parent


def report_day(value):
    match = re.match(r'^(\d{4})[.-](\d{2})[.-](\d{2})(?:$|\s)', str(value or ''))
    if not match:
        raise ValueError('Тайлант огноо YYYY-MM-DD эсвэл хуучин YYYY.MM.DD - ... хэлбэртэй байна.')
    return date(*map(int, match.groups())).isoformat()


def percentage(value):
    text = str(value or '').strip().rstrip('%').strip()
    try:
        number = float(text)
    except ValueError:
        return None
    return number if 0 <= number <= 100 else None


def init_db(path, admin_password=None):
    """Idempotent migration; backup existing DB before schema changes, never delete reports."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    exists = path.exists()
    with closing(sqlite3.connect(path)) as db, db:
        version = db.execute('PRAGMA user_version').fetchone()[0]
        if exists and version < 2:
            backup = path.with_name(path.name + '.before-v2-' + datetime.now().strftime('%Y%m%d%H%M%S%f'))
            with closing(sqlite3.connect(backup)) as dest:
                db.backup(dest)
        db.execute('CREATE TABLE IF NOT EXISTS departments (name TEXT PRIMARY KEY)')
        db.executemany('INSERT OR IGNORE INTO departments VALUES (?)', [(d,) for d in DEPARTMENTS])
        db.execute('CREATE TABLE IF NOT EXISTS users (id INTEGER PRIMARY KEY AUTOINCREMENT, email TEXT UNIQUE NOT NULL, password TEXT NOT NULL, role TEXT NOT NULL DEFAULT \'user\', department TEXT)')
        if 'department' not in {r[1] for r in db.execute('PRAGMA table_info(users)')}:
            db.execute('ALTER TABLE users ADD COLUMN department TEXT')
        db.execute('CREATE TABLE IF NOT EXISTS reports (id INTEGER PRIMARY KEY AUTOINCREMENT, department TEXT NOT NULL, report_date TEXT NOT NULL, func_area TEXT NOT NULL, task_desc TEXT NOT NULL, task_result TEXT, progress TEXT NOT NULL, assignee TEXT, status TEXT DEFAULT \'submitted\', created_at TEXT)')
        columns = {r[1] for r in db.execute('PRAGMA table_info(reports)')}
        for name, kind in [('issue', "TEXT NOT NULL DEFAULT ''"), ('report_day', 'TEXT'), ('created_at', 'TEXT')]:
            if name not in columns:
                db.execute(f'ALTER TABLE reports ADD COLUMN {name} {kind}')
        for ident, value in db.execute('SELECT id, report_date FROM reports WHERE report_day IS NULL').fetchall():
            try:
                day = report_day(value)
            except ValueError:
                continue  # Ambiguous legacy dates remain visible in the journal, not invented.
            db.execute('UPDATE reports SET report_day=? WHERE id=?', (day, ident))
        db.execute('CREATE TABLE IF NOT EXISTS categories (id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT UNIQUE NOT NULL)')
        if version < 2:
            cols = {r[1] for r in db.execute('PRAGMA table_info(categories)')}
            old = db.execute('SELECT id,name' + (',department' if 'department' in cols else '') + ' FROM categories').fetchall()
            db.execute('ALTER TABLE categories RENAME TO categories_legacy_v1')
            db.execute('CREATE TABLE categories (id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT NOT NULL, department TEXT REFERENCES departments(name), UNIQUE(department,name))')
            for row in old:
                ident, name = row[:2]
                dept = row[2] if len(row) > 2 and row[2] in DEPARTMENTS else None
                canonical = 'Архив, албан хэрэг хөтлөлт' if name == 'Архив, албан хэрэг' else name
                matches = [d for d, names in FUNCTIONS.items() if canonical in names]
                if not dept and len(matches) == 1:
                    dept = matches[0]
                # Unknown mappings are retained but excluded from employee dropdowns.
                db.execute('INSERT OR IGNORE INTO categories(id,name,department) VALUES (?,?,?)', (ident, name, dept))
            for dept, names in FUNCTIONS.items():
                db.executemany('INSERT OR IGNORE INTO categories(name,department) VALUES (?,?)', [(n, dept) for n in names])
            db.execute('PRAGMA user_version=2')
        db.execute('CREATE INDEX IF NOT EXISTS report_scope ON reports(department,report_day)')
        db.execute("CREATE TABLE IF NOT EXISTS report_revisions (id INTEGER PRIMARY KEY AUTOINCREMENT, report_id INTEGER NOT NULL REFERENCES reports(id), task_result TEXT NOT NULL, submitted_by TEXT NOT NULL, submitted_at TEXT NOT NULL, status TEXT NOT NULL DEFAULT 'pending', reviewed_by TEXT, reviewed_at TEXT, review_note TEXT, previous_progress TEXT, previous_result TEXT)")
        db.execute("CREATE UNIQUE INDEX IF NOT EXISTS pending_revision ON report_revisions(report_id) WHERE status='pending'")
        if not db.execute('SELECT 1 FROM users').fetchone() and admin_password:
            db.execute('INSERT INTO users(email,password,role,department) VALUES (?,?,?,?)', ('admin@uag.mn', generate_password_hash(admin_password), 'admin', 'Бүх хэлтэс'))


def create_app(test_config=None):
    app = Flask(__name__)
    secret_path = ROOT / '.secret_key'
    if not secret_path.exists():
        try:
            with secret_path.open('x', encoding='ascii') as out:
                out.write(secrets.token_hex(32))
        except FileExistsError:
            pass
    app.config.update(SECRET_KEY=os.getenv('UAG_SECRET_KEY') or secret_path.read_text().strip(),
                      DATABASE=os.getenv('UAG_DATABASE', str(ROOT / 'uag_reports_system.db')),
                      MAX_CONTENT_LENGTH=8 * 1024 * 1024,
                      SESSION_COOKIE_HTTPONLY=True, SESSION_COOKIE_SAMESITE='Lax',
                      SESSION_COOKIE_SECURE=os.getenv('UAG_HTTPS') == '1')
    if test_config:
        app.config.update(test_config)
    init_db(app.config['DATABASE'], os.getenv('UAG_ADMIN_PASSWORD'))

    def db():
        if 'db' not in g:
            g.db = sqlite3.connect(app.config['DATABASE'])
            g.db.row_factory = sqlite3.Row
            g.db.execute('PRAGMA foreign_keys=ON')
        return g.db

    @app.teardown_appcontext
    def close_db(error=None):
        if 'db' in g:
            g.db.close()

    @app.before_request
    def authenticate():
        g.user = None
        if session.get('user'):
            g.user = db().execute('SELECT * FROM users WHERE email=?', (session['user'],)).fetchone()
        if g.user:
            # Reload authorization from DB on every request, including downloads.
            if g.user['role'] not in ROLES or (g.user['role'] not in LEADERS and g.user['department'] not in DEPARTMENTS):
                abort(403, 'Хэрэглэгчийн хэлтэс эсвэл эрхийн тохиргоог админ засах шаардлагатай.')
            session.update(role=g.user['role'], department=g.user['department'])
        if request.method == 'POST':
            supplied = request.form.get('csrf_token', '')
            if not supplied or not secrets.compare_digest(supplied, session.get('csrf_token', '')):
                abort(400, 'Хуудас шинэчлээд дахин оролдоно уу (CSRF).')

    @app.context_processor
    def context():
        session.setdefault('csrf_token', secrets.token_hex(24))
        return dict(user=g.user, leader=bool(g.user and g.user['role'] in LEADERS), departments=DEPARTMENTS,
                    csrf_token=session['csrf_token'])

    def required(leader=False):
        def decorate(f):
            @wraps(f)
            def call(*args, **kwargs):
                if not g.user:
                    return redirect(url_for('login'))
                if leader and g.user['role'] not in LEADERS:
                    abort(403)
                return f(*args, **kwargs)
            return call
        return decorate

    def scope(value=None):
        if g.user['role'] not in LEADERS:
            return g.user['department']
        value = value if value is not None else request.args.get('dept', '')
        if value and value not in DEPARTMENTS:
            abort(400, 'Хэлтэс буруу байна.')
        return value

    def dates(default=False):
        start, end = request.args.get('start_date', ''), request.args.get('end_date', '')
        today = date.today()
        preset = request.args.get('preset')
        if preset in ('this_week', 'last_week') or (default and not start and not end):
            monday = today - timedelta(days=today.weekday() + (7 if preset != 'this_week' else 0))
            start, end = monday.isoformat(), (monday + timedelta(days=6)).isoformat()
        elif preset == 'this_month':
            start, end = today.replace(day=1).isoformat(), today.isoformat()
        try:
            if start: date.fromisoformat(start)
            if end: date.fromisoformat(end)
        except ValueError:
            abort(400, 'Огноо буруу байна.')
        if start and end and start > end:
            abort(400, 'Эхлэх огноо дуусах огнооноос хойш байна.')
        return start, end

    def rows(dept, start='', end='', q=''):
        conditions, params = [], []
        if dept:
            conditions.append('department=?'); params.append(dept)
        if start:
            conditions.append('report_day>=?'); params.append(start)
        if end:
            conditions.append('report_day<=?'); params.append(end)
        if q:
            conditions.append('(task_desc LIKE ? OR assignee LIKE ? OR task_result LIKE ?)')
            params.extend(['%' + q + '%'] * 3)
        return [dict(r) for r in db().execute('SELECT * FROM reports' + (' WHERE ' + ' AND '.join(conditions) if conditions else '') + ' ORDER BY report_day DESC,id DESC', params)]

    def categories(dept):
        return [r[0] for r in db().execute('SELECT name FROM categories WHERE department=? ORDER BY id', (dept,))]

    @app.route('/login', methods=['GET', 'POST'])
    def login():
        error = None
        if request.method == 'POST':
            account = db().execute('SELECT * FROM users WHERE email=?', (request.form.get('email', '').strip(),)).fetchone()
            if account and check_password_hash(account['password'], request.form.get('password', '')):
                session.clear()
                session['user'] = account['email']
                return redirect(url_for('dashboard'))
            error = 'Имэйл эсвэл нууц үг буруу байна.'
        return render_template('login.html', error=error)

    @app.post('/logout')
    def logout():
        session.clear()
        return redirect(url_for('login'))

    @app.get('/dashboard')
    @required()
    def dashboard():
        from reporting import metrics
        dept, (start, end) = scope(), dates(default=True)
        records = rows(dept, start, end)
        selected = [dept] if dept else DEPARTMENTS
        breakdown = [dict(name=d, **metrics([r for r in records if r['department'] == d])) for d in selected]
        return render_template('dashboard.html', stats=metrics(records), breakdown=breakdown,
                               recent=records[:6], issues=[r for r in records if r.get('issue')][:5],
                               dept=dept, start=start, end=end)

    @app.get('/')
    @required()
    def index():
        dept = scope()
        entry_dept = dept or (g.user['department'] if g.user['department'] in DEPARTMENTS else DEPARTMENTS[0])
        start, end = dates()
        records = rows(dept, start, end, request.args.get('q', ''))
        incomplete = request.args.get('incomplete') == '1'
        for record in records:
            record['is_incomplete'] = percentage(record['progress']) != 100
            record['pending_revision'] = db().execute("SELECT id FROM report_revisions WHERE report_id=? AND status='pending'", (record['id'],)).fetchone()
        if incomplete:
            records = [r for r in records if r['is_incomplete']]
        return render_template('index.html', rows=records, incomplete=incomplete,
                               dept=dept, entry_dept=entry_dept, categories=categories(entry_dept),
                               start=start, end=end, today=date.today().isoformat())

    @app.route('/complete/<int:id>', methods=['GET', 'POST'])
    @required()
    def complete_report(id):
        dept = scope('')
        report = db().execute('SELECT * FROM reports WHERE id=?' + (' AND department=?' if dept else ''), (id,dept) if dept else (id,)).fetchone()
        if not report: abort(404)
        if percentage(report['progress']) == 100: abort(409, 'Ажил аль хэдийн 100% гүйцэтгэлтэй байна.')
        pending = db().execute("SELECT id FROM report_revisions WHERE report_id=? AND status='pending'",(id,)).fetchone()
        if pending: abort(409, 'Таны нөхөлтийг удирдлага хянаж байна. Давхар илгээх шаардлагагүй.')
        history = db().execute('SELECT * FROM report_revisions WHERE report_id=? ORDER BY id DESC',(id,)).fetchall()
        if request.method == 'POST':
            result = request.form.get('task_result','').strip()
            if not result or len(result)>12000: abort(400, 'Гүйцэтгэсэн ажлын үр дүнг 1–12000 тэмдэгтээр оруулна уу.')
            try:
                with db():
                    # Lock before checking status to serialize completion and approval.
                    db().execute('BEGIN IMMEDIATE')
                    latest = db().execute('SELECT * FROM reports WHERE id=?',(id,)).fetchone()
                    if percentage(latest['progress']) == 100: abort(409, 'Ажил аль хэдийн дууссан.')
                    needs_review = latest['status'] == 'approved'
                    db().execute('INSERT INTO report_revisions(report_id,task_result,submitted_by,submitted_at,status,previous_progress,previous_result) VALUES (?,?,?,?,?,?,?)',
                                 (id,result,g.user['email'],datetime.now().isoformat(timespec='seconds'),'pending' if needs_review else 'applied',latest['progress'],latest['task_result']))
                    if not needs_review:
                        db().execute("UPDATE reports SET task_result=?,progress='100%',status='submitted' WHERE id=?",(result,id))
            except sqlite3.IntegrityError: abort(409, 'Энэ ажлын нөхөлт аль хэдийн хянагдаж байна.')
            return redirect(url_for('index',incomplete='1' if needs_review else '',q=report['task_desc'][:50]))
        return render_template('complete.html',report=report,history=history)

    @app.get('/reviews')
    @required(leader=True)
    def reviews():
        revisions = db().execute("SELECT v.*,r.department,r.task_desc,r.progress,r.task_result AS current_result FROM report_revisions v JOIN reports r ON r.id=v.report_id WHERE v.status='pending' ORDER BY v.id").fetchall()
        return render_template('reviews.html',revisions=revisions)

    @app.post('/reviews/<int:id>')
    @required(leader=True)
    def review_revision(id):
        action=request.form.get('action')
        if action not in ('approve','return'): abort(400)
        note=request.form.get('note','').strip()
        if action=='return' and not note: abort(400,'Буцаах шалтгаанаа бичнэ үү.')
        with db():
            db().execute('BEGIN IMMEDIATE')
            revision=db().execute("SELECT * FROM report_revisions WHERE id=? AND status='pending'",(id,)).fetchone()
            if not revision: abort(409,'Хүсэлтийг өмнө нь шийдвэрлэсэн байна.')
            if action=='approve':
                db().execute("UPDATE reports SET task_result=?,progress='100%',status='approved' WHERE id=?",(revision['task_result'],revision['report_id']))
            db().execute('UPDATE report_revisions SET status=?,reviewed_by=?,reviewed_at=?,review_note=? WHERE id=?',
                         ('approved' if action=='approve' else 'returned',g.user['email'],datetime.now().isoformat(timespec='seconds'),note,id))
        return redirect(url_for('reviews'))

    def validated_report(dept, form, i=None):
        def field(name, default=''):
            if i is None: return form.get(name, default).strip()
            values = form.getlist(name)
            return values[i].strip() if i < len(values) else default
        func = field('func_area')
        if func not in categories(dept):
            abort(400, 'Сонгосон чиг үүрэг тухайн хэлтэст хамаарахгүй байна.')
        desc = field('task_desc')
        if not desc: abort(400, 'Ажлын мэдээлэл хоосон байна.')
        progress = percentage(field('progress'))
        if progress is None: abort(400, 'Явц 0-100 хооронд тоо байна.')
        try:
            day = report_day(form.get('report_date'))
        except ValueError as ex:
            abort(400, str(ex))
        values = [day, func, desc, field('task_result'), f'{progress:g}%', field('assignee'), field('issue'), day]
        if any(len(v) > 12000 for v in values): abort(400, 'Нэг талбар 12000 тэмдэгтээс хэтэрсэн байна.')
        return values

    @app.post('/save')
    @required()
    def save():
        dept = scope(request.form.get('department', ''))
        if not dept: abort(400, 'Тайлан оруулах хэлтсээ сонгоно уу.')
        count = len(request.form.getlist('func_area'))
        if not 1 <= count <= 500: abort(400, 'Мөрийн тоо 1-500 байна.')
        if any(len(request.form.getlist(k)) != count for k in ('task_desc', 'task_result', 'progress', 'assignee', 'issue')):
            abort(400, 'Тайлангийн мөрийн талбарууд дутуу байна.')
        batch = [validated_report(dept, request.form, i) for i in range(count)]
        with db():
            db().executemany('INSERT INTO reports(department,report_date,func_area,task_desc,task_result,progress,assignee,issue,report_day,status,created_at) VALUES (?,?,?,?,?,?,?,?,?,\'submitted\',?)',
                             [(dept, *v, datetime.now().isoformat(timespec='seconds')) for v in batch])
        return redirect(url_for('index', dept=dept))

    @app.route('/edit_report/<int:id>', methods=['GET', 'POST'])
    @required()
    def edit_report(id):
        dept = scope('')
        report = db().execute('SELECT * FROM reports WHERE id=?' + (' AND department=?' if dept else ''), (id, dept) if dept else (id,)).fetchone()
        if not report: abort(404)
        if report['status'] == 'approved': abort(409, 'Батлагдсан тайлан. Удирдлага буцаасны дараа засна.')
        if db().execute("SELECT 1 FROM report_revisions WHERE report_id=? AND status='pending'",(id,)).fetchone():
            abort(409, 'Нөхөлт хянагдаж байгаа тул энгийн засвар түр хаалттай.')
        if request.method == 'POST':
            values = validated_report(report['department'], request.form)
            with db():
                changed = db().execute('UPDATE reports SET report_date=?,func_area=?,task_desc=?,task_result=?,progress=?,assignee=?,issue=?,report_day=? WHERE id=? AND status!=\'approved\'', (*values, id))
                if not changed.rowcount: abort(409)
            return redirect(url_for('index'))
        return render_template('edit.html', report=report, categories=categories(report['department']))

    @app.get('/analytics')
    @app.get('/reports')
    @required()
    def report_page():
        from reporting import build_report
        dept, (start, end) = scope(), dates(default=True)
        model = build_report(rows(dept, start, end), rows(dept, '', end), dept, start, end, g.user['role'] in LEADERS)
        query = dict(dept=dept, start_date=start, end_date=end)
        return render_template('report.html', model=model, dept=dept, start=start, end=end, query=query)

    @app.get('/reports/export/<kind>')
    @required()
    def export_report(kind):
        from reporting import build_report, export_pdf, export_pptx, ReportOverflow
        if kind not in ('pdf', 'pptx'): abort(404)
        dept, (start, end) = scope(), dates(default=True)
        model = build_report(rows(dept, start, end), rows(dept, '', end), dept, start, end, g.user['role'] in LEADERS)
        try:
            content = export_pdf(model) if kind == 'pdf' else export_pptx(model)
        except ReportOverflow as exc:
            abort(422, str(exc))
        return send_file(io.BytesIO(content), as_attachment=True, download_name=f'UAG_{start}_{end}.{kind}',
                         mimetype='application/pdf' if kind == 'pdf' else 'application/vnd.openxmlformats-officedocument.presentationml.presentation')

    @app.get('/admin')
    @required(leader=True)
    def admin():
        page = request.args.get('page', 'categories')
        if page in ('reports', 'analytics', 'ai'):
            return redirect(url_for('report_page', **{k: v for k, v in request.args.items() if k != 'page'}))
        return render_template('admin.html', page=page, categories=db().execute('SELECT * FROM categories ORDER BY department,id').fetchall(),
                               users=db().execute('SELECT id,email,role,department FROM users').fetchall(), roles=sorted(ROLES))

    @app.post('/admin/add_category')
    @required(leader=True)
    def add_category():
        dept, name = request.form.get('department'), request.form.get('category_name', '').strip()
        if dept not in DEPARTMENTS or not name or len(name) > 150: abort(400)
        with db():
            db().execute('INSERT OR IGNORE INTO categories(name,department) VALUES (?,?)', (name, dept))
        return redirect(url_for('admin'))

    @app.post('/admin/map_category/<int:id>')
    @required(leader=True)
    def map_category(id):
        dept = request.form.get('department')
        if dept not in DEPARTMENTS: abort(400)
        try:
            with db(): db().execute('UPDATE categories SET department=? WHERE id=?', (dept, id))
        except sqlite3.IntegrityError: abort(409, 'Энэ хэлтэст ижил нэртэй чиг үүрэг байна.')
        return redirect(url_for('admin'))

    @app.post('/admin/delete_category/<int:id>')
    @required(leader=True)
    def delete_category(id):
        with db(): db().execute('DELETE FROM categories WHERE id=?', (id,))
        return redirect(url_for('admin'))

    @app.post('/admin/add_user')
    @required(leader=True)
    def add_user():
        if g.user['role'] != 'admin': abort(403)
        email, password, role, dept = (request.form.get(k, '').strip() for k in ('email', 'password', 'role', 'department'))
        if '@' not in email or len(password) < 8 or role not in ROLES or dept not in DEPARTMENTS: abort(400, 'Имэйл, хэлтэс, эрхээ шалгана уу. Нууц үг 8-аас доошгүй тэмдэгт байна.')
        try:
            with db(): db().execute('INSERT INTO users(email,password,role,department) VALUES (?,?,?,?)', (email, generate_password_hash(password), role, dept))
        except sqlite3.IntegrityError: abort(409, 'Имэйл бүртгэлтэй байна.')
        return redirect(url_for('admin', page='others'))

    @app.post('/admin/delete_user/<int:id>')
    @required(leader=True)
    def delete_user(id):
        if g.user['role'] != 'admin' or id == g.user['id']: abort(403)
        with db(): db().execute('DELETE FROM users WHERE id=?', (id,))
        return redirect(url_for('admin', page='others'))

    @app.post('/admin/toggle_report_status/<int:id>')
    @required(leader=True)
    def toggle_report_status(id):
        if db().execute("SELECT 1 FROM report_revisions WHERE report_id=? AND status='pending'",(id,)).fetchone():
            abort(409, 'Энэ ажлын нөхөлтийг эхлээд Нөхөлт хянах цэснээс шийдвэрлэнэ үү.')
        with db():
            db().execute("UPDATE reports SET status=CASE WHEN status='approved' THEN 'submitted' ELSE 'approved' END WHERE id=?", (id,))
        return redirect(url_for('index'))

    @app.post('/admin/delete_report/<int:id>')
    @required(leader=True)
    def delete_report(id):
        if db().execute('SELECT 1 FROM report_revisions WHERE report_id=?',(id,)).fetchone():
            abort(409, 'Нөхөлтийн түүхтэй ажлыг устгахгүй. Түүхийг хадгална.')
        with db(): db().execute('DELETE FROM reports WHERE id=?', (id,))
        return redirect(url_for('index'))

    @app.get('/admin/export_csv')
    @app.get('/export_csv')
    @required()
    def export_csv():
        dept, (start, end) = scope(), dates()
        output = io.StringIO(newline=''); output.write('\ufeff')
        writer = csv.writer(output)
        writer.writerow(['Тайлант огноо', 'Хэлтэс', 'Чиг үүрэг', 'Ажил', 'Үр дүн', 'Явц', 'Хариуцагч', 'Шийдвэрлүүлэх асуудал', 'Төлөв'])
        for r in rows(dept, start, end, request.args.get('q', '')):
            values = [r[k] or '' for k in ('report_date', 'department', 'func_area', 'task_desc', 'task_result', 'progress', 'assignee', 'issue', 'status')]
            writer.writerow(["'" + v if v.lstrip().startswith(('=', '+', '-', '@', '\t', '\r')) else v for v in values])
        return send_file(io.BytesIO(output.getvalue().encode('utf-8')), as_attachment=True, download_name='UAG_reports.csv', mimetype='text/csv')

    @app.get('/template.xlsx')
    @required()
    def excel_template():
        # User-facing import format is generated at runtime, not a reference workbook modification.
        from openpyxl import Workbook
        dept = scope() or DEPARTMENTS[0]
        workbook = Workbook(); sheet = workbook.active; sheet.title = 'Тайлан'
        sheet.append(['Чиг үүрэг', 'Ажлын мэдээлэл', 'Үр дүн', 'Явц', 'Хариуцагч', 'Шийдвэрлүүлэх асуудал'])
        for name in categories(dept): sheet.append([name, '', '', '', '', ''])
        data = io.BytesIO(); workbook.save(data); data.seek(0)
        return send_file(data, as_attachment=True, download_name='UAG_template.xlsx')

    @app.post('/import')
    @required()
    def import_excel():
        from openpyxl import load_workbook
        dept = scope(request.form.get('department', ''))
        if not dept: abort(400)
        upload = request.files.get('file')
        if not upload: abort(400)
        try:
            workbook = load_workbook(upload, read_only=True, data_only=True)
            sheet = workbook.worksheets[0]
            if sheet.max_row > 501: abort(400, 'Импорт 500 хүртэл мөртэй байна.')
            values = list(sheet.iter_rows(values_only=True))
        except Exception as ex:
            from werkzeug.exceptions import HTTPException
            if isinstance(ex, HTTPException): raise
            abort(400, 'XLSX файлыг унших боломжгүй байна.')
        batch = []
        for row in values[1:]:
            if len(row) < 2 or not row[1]: continue
            row = list(row) + [''] * 6
            form = dict(zip(('func_area', 'task_desc', 'task_result', 'progress', 'assignee', 'issue'), [str(v) if v is not None else '' for v in row[:6]]))
            form['report_date'] = request.form.get('report_date', '')
            batch.append(validated_report(dept, form))
        if not batch: abort(400, 'Импортлох ажил алга. Системийн загварыг ашиглана уу.')
        with db():
            db().executemany('INSERT INTO reports(department,report_date,func_area,task_desc,task_result,progress,assignee,issue,report_day,status,created_at) VALUES (?,?,?,?,?,?,?,?,?,\'submitted\',?)', [(dept, *v, datetime.now().isoformat(timespec='seconds')) for v in batch])
        return redirect(url_for('index', dept=dept))

    @app.errorhandler(400)
    @app.errorhandler(403)
    @app.errorhandler(404)
    @app.errorhandler(409)
    @app.errorhandler(422)
    def error(exc):
        return render_template('error.html', error=exc), exc.code

    return app


if __name__ == '__main__':
    application = create_app()
    print('UAG: http://127.0.0.1:5000')
    application.run(host='127.0.0.1', port=5000, debug=False)
